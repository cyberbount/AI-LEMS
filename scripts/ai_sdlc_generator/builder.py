# -*- coding: utf-8 -*-
"""
Trình tạo Báo cáo Thực hành AI-Augmented SDLC Master cho AI-LEMS (Đề tài 23).
Tạo song song:
  - docs/Bao_Cao_Thuc_Hanh_AI_Augmented_SDLC_AI_LEMS.md
  - docs/Bao_Cao_Thuc_Hanh_AI_Augmented_SDLC_AI_LEMS.docx
"""

import os
import sys
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from .part_frontmatter import get_frontmatter_markdown
from .part_requirements import get_requirements_markdown
from .part_architecture_db import get_architecture_db_markdown
from .part_core_review import get_core_review_markdown
from .part_ai_assistant import get_ai_assistant_markdown
from .part_final_appendix import get_final_appendix_markdown

BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DOCS_DIR = os.path.join(BASE_DIR, "docs")
MD_OUTPUT = os.path.join(DOCS_DIR, "Bao_Cao_Thuc_Hanh_AI_Augmented_SDLC_AI_LEMS.md")
DOCX_OUTPUT = os.path.join(DOCS_DIR, "Bao_Cao_Thuc_Hanh_AI_Augmented_SDLC_AI_LEMS.docx")

# Màu sắc
NAVY = RGBColor(0x1A, 0x36, 0x5D)
DARK = RGBColor(0x2D, 0x37, 0x48)
BLUE = RGBColor(0x2B, 0x6C, 0xB0)
GREEN = RGBColor(0x22, 0x86, 0x3A)
RED = RGBColor(0xCB, 0x24, 0x31)
GREY = RGBColor(0x71, 0x80, 0x96)

def set_cell_bg(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def build_docx_from_md(md_text, docx_path):
    doc = Document()
    
    # Cấu hình Margins
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5) # lề đóng gáy
        section.right_margin = Cm(2.0)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing = 1.2
    style.paragraph_format.space_after = Pt(4)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    # ================= TRANG BÌA =================
    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(36)
    
    r = cover.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO — TRƯỜNG ĐẠI HỌC BÁCH KHOA\nKHOA CÔNG NGHỆ THÔNG TIN\n\n\n")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = GREY

    r = cover.add_run("BÁO CÁO THỰC HÀNH TỔNG HỢP\nXÂY DỰNG HỆ THỐNG QUẢN LÝ THIẾT BỊ PHÒNG LAB ĐIỆN TỬ - IOT CÓ TÍCH HỢP AI\n")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = NAVY

    r = cover.add_run("QUY TRÌNH AI-AUGMENTED SDLC VỚI CODEX / ANTIGRAVITY\n\n")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = BLUE

    r = cover.add_run("Đề tài 23: Hệ thống AI-LEMS (AI-Augmented Lab Equipment Management System)\n\n\n\n")
    r.italic = True
    r.font.size = Pt(11.5)

    # Bảng thông tin trang bìa
    t_info = doc.add_table(rows=0, cols=2)
    t_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_info.style = "Table Grid"
    info_rows = [
        ("Môn học:", "Ứng dụng Trí tuệ Nhân tạo trong Phát triển Phần mềm"),
        ("Giảng viên hướng dẫn:", "TS. Nguyễn Đình Dũng"),
        ("Nhóm thực hiện:", "Nhóm nghiên cứu Đề tài 23 (AI-LEMS)"),
        ("Quy trình chuẩn hóa:", "AI-Augmented SDLC (Skills, Tools, MCP, Human Gates)"),
        ("Mô hình AI tích hợp:", "Local LLM Ollama Qwen 2.5 3B (Quantized Q4_K_M)"),
        ("Thời gian thực hiện:", "Tháng 09/2026 - Tháng 10/2026"),
        ("Trạng thái nghiệm thu:", "100% HOÀN TẤT & ĐẠT CHUẨN XUẤT SẮC"),
    ]
    for k, v in info_rows:
        row = t_info.add_row().cells
        row[0].width = Cm(5.5)
        row[1].width = Cm(11.0)
        set_cell_bg(row[0], "F2F4F8")
        set_cell_margins(row[0], 60, 60, 100, 100)
        set_cell_margins(row[1], 60, 60, 100, 100)
        pk = row[0].paragraphs[0]
        pk.paragraph_format.space_after = Pt(2)
        rk = pk.add_run(k)
        rk.bold = True
        rk.font.size = Pt(10)
        pv = row[1].paragraphs[0]
        pv.paragraph_format.space_after = Pt(2)
        rv = pv.add_run(v)
        rv.font.size = Pt(10)
        if "100%" in v:
            rv.bold = True
            rv.font.color.rgb = GREEN

    doc.add_page_break()

    # ================= PARSER NỘI DUNG MARKDOWN SANG DOCX =================
    lines = md_text.split("\n")
    idx = 0
    in_code_block = False
    code_lang = ""
    code_lines = []

    in_table = False
    table_rows = []

    while idx < len(lines):
        line = lines[idx]
        stripped = line.strip()

        # Xử lý Code block
        if stripped.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lang = stripped[3:].strip()
                code_lines = []
                idx += 1
                continue
            else:
                in_code_block = False
                code_content = "\n".join(code_lines)
                
                # Render code block vào docx
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_bg(cell, "F7FAFC")
                set_cell_margins(cell, 80, 80, 120, 120)
                cell.width = Cm(16.5)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                
                # Tiêu đề code block
                if code_lang:
                    r_lang = p.add_run(f"[{code_lang.upper()}]\n")
                    r_lang.bold = True
                    r_lang.font.name = "Times New Roman"
                    r_lang.font.size = Pt(8.5)
                    r_lang.font.color.rgb = BLUE

                r_code = p.add_run(code_content)
                r_code.font.name = "Courier New"
                r_code.font.size = Pt(8.5)
                r_code.font.color.rgb = DARK
                
                doc.add_paragraph().paragraph_format.space_after = Pt(3)
                idx += 1
                continue

        if in_code_block:
            code_lines.append(line)
            idx += 1
            continue

        # Xử lý Table Markdown
        if stripped.startswith("|") and stripped.endswith("|"):
            if not in_table:
                in_table = True
                table_rows = [stripped]
            else:
                table_rows.append(stripped)
            idx += 1
            continue
        elif in_table:
            # Kết thúc bảng, render bảng
            in_table = False
            # Lọc bỏ dòng separator |---|---|
            parsed_rows = []
            for r_str in table_rows:
                cells = [c.strip() for c in r_str.strip("|").split("|")]
                if all(re.match(r"^:?-+:?$", c) for c in cells if c):
                    continue
                parsed_rows.append(cells)
            
            if parsed_rows:
                num_cols = max(len(r) for r in parsed_rows)
                t = doc.add_table(rows=0, cols=num_cols)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                t.style = "Table Grid"
                for r_i, r_data in enumerate(parsed_rows):
                    row_cells = t.add_row().cells
                    for c_i in range(num_cols):
                        val = r_data[c_i] if c_i < len(r_data) else ""
                        set_cell_margins(row_cells[c_i], 60, 60, 80, 80)
                        p = row_cells[c_i].paragraphs[0]
                        p.paragraph_format.space_after = Pt(2)
                        
                        # Làm sạch bold tags
                        clean_val = val.replace("**", "").replace("*", "")
                        r = p.add_run(clean_val)
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(9)
                        
                        if r_i == 0:
                            set_cell_bg(row_cells[c_i], "1A365D")
                            r.bold = True
                            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        else:
                            if c_i == 0:
                                set_cell_bg(row_cells[c_i], "F2F4F8")
                                r.bold = True
                            if "PASS" in val or "ĐẠT" in val or "APPROVED" in val:
                                r.bold = True
                                r.font.color.rgb = GREEN
                            elif "FAIL" in val or "HIGH" in val or "LỖ HỔNG" in val:
                                r.bold = True
                                r.font.color.rgb = RED
                doc.add_paragraph().paragraph_format.space_after = Pt(3)

        # Xử lý Headings
        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(stripped[2:])
            r.bold = True
            r.font.size = Pt(15)
            r.font.color.rgb = NAVY
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(stripped[3:])
            r.bold = True
            r.font.size = Pt(13)
            r.font.color.rgb = NAVY
        elif stripped.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(stripped[4:])
            r.bold = True
            r.font.size = Pt(11.5)
            r.font.color.rgb = BLUE
        elif stripped.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(stripped[5:])
            r.bold = True
            r.font.size = Pt(11)
            r.font.color.rgb = DARK
        elif stripped.startswith("> "):
            # Callout / Blockquote
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            set_cell_bg(cell, "EEF6FF")
            set_cell_margins(cell, 80, 80, 120, 120)
            cell.width = Cm(16.5)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(stripped[2:].replace("**", ""))
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.italic = True
            r.font.color.rgb = NAVY
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(stripped[2:].replace("**", ""))
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
        elif re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            txt_num = re.sub(r"^\d+\.\s", "", stripped).replace("**", "")
            r = p.add_run(txt_num)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
        elif stripped == "---":
            # Phân cách phần
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run("_________________________________________________________________________________")
            r.font.size = Pt(8)
            r.font.color.rgb = GREY
        elif stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(stripped.replace("**", ""))
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)

        idx += 1

    doc.save(docx_path)
    print(f"Đã xuất thành công tệp Word: {docx_path}")

def run_build():
    print("Đang tổng hợp toàn bộ các phân hệ tài liệu...")
    
    sections = [
        get_frontmatter_markdown(),
        get_requirements_markdown(),
        get_architecture_db_markdown(),
        get_core_review_markdown(),
        get_ai_assistant_markdown(),
        get_final_appendix_markdown(),
    ]
    
    full_markdown = "\n\n".join(sections)
    
    # Ghi file Markdown
    with open(MD_OUTPUT, "w", encoding="utf-8") as f:
        f.write(full_markdown)
    print(f"Đã xuất thành công tệp Markdown: {MD_OUTPUT} (Kích thước: {len(full_markdown):,} ký tự)")
    
    # Ghi file Word DOCX
    build_docx_from_md(full_markdown, DOCX_OUTPUT)
    print("HOÀN TẤT XUẤT BẢN TOÀN BỘ TÀI LIỆU AI-AUGMENTED SDLC!")

if __name__ == "__main__":
    run_build()
